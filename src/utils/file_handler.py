"""
语料文件工具：MD5、按后缀递归列举、用 LangChain Loader 读成 Document。

供 VectorStoreService 入库使用。MD5 为遗留辅助；主路径增量去重用 content_hash。
"""
from __future__ import annotations

import hashlib
import os

from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document

from utils.logger_handler import logger


def get_file_md5_hex(filepath: str) -> str | None:
    """计算文件 MD5 十六进制串；路径不是普通文件时返回 None。"""
    if not os.path.isfile(filepath):
        logger.error("[md5计算]不是有效文件：%s", filepath)
        return None

    md5_obj = hashlib.md5()
    chunk_size = 4096  # 分片避免大文件占满内存
    try:
        with open(filepath, "rb") as f:
            while chunk := f.read(chunk_size):
                md5_obj.update(chunk)
        return md5_obj.hexdigest()
    except Exception as e:
        logger.error("计算文件 %s 的 md5 失败：%s", filepath, e)
        return None


def listdir_with_allowed_type(path: str, allowed_type: tuple[str, ...]):
    """递归列出目录中后缀匹配的文件（allowed_type 如 txt/pdf/md，可带或不带点）。"""
    files: list[str] = []
    if not os.path.isdir(path):
        logger.error("[listdir_with_allowed_type]%s 不是文件夹", path)
        return tuple()

    suffixes = []
    for s in allowed_type:
        s = str(s).lower().strip()
        if not s:
            continue
        if not s.startswith("."):
            s = "." + s
        suffixes.append(s)
    suffix_tuple = tuple(suffixes)
    if not suffix_tuple:
        return tuple()

    for root, _dirs, names in os.walk(path):
        for name in names:
            if name.lower().endswith(suffix_tuple):
                files.append(os.path.join(root, name))
    return tuple(files)


def pdf_loader(filepath: str, password: str | None = None) -> list[Document]:
    """加载 PDF 为 LangChain Document 列表。"""
    if password:
        return PyPDFLoader(filepath, password=password).load()
    return PyPDFLoader(filepath).load()


def txt_loader(filepath: str) -> list[Document]:
    """按 UTF-8 加载纯文本。"""
    return TextLoader(filepath, encoding="utf-8").load()


def md_loader(filepath: str) -> list[Document]:
    """Markdown 与 txt 共用 TextLoader（切分阶段再用 MarkdownHeaderTextSplitter）。"""
    return TextLoader(filepath, encoding="utf-8").load()


def docx_loader(filepath: str) -> list[Document]:
    """加载 Word 文档：段落 + 表格按行拼接为纯文本。"""
    from docx import Document as DocxDocument

    doc = DocxDocument(filepath)
    parts: list[str] = []
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if (c.text or "").strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts)
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": filepath})]
